#!/usr/bin/env python3
"""
Insert a summarized itinerary section into a styled .docx base file.
Inserts before the "Trip at a Glance" heading, preserving all existing formatting.

Usage:
    python3 src/tour-planner/insert_summary.py \
        --base output/trip-itinerary-v2.docx \
        --summary output/trip-itinerary-latest.md \
        --output output/trip-itinerary-latest.docx
"""

import argparse
import re
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def parse_summary_from_md(md_path: Path) -> list:
    """Extract the Summarized Itinerary section from the .md file.
    Returns a list of (text, type) tuples."""

    lines = md_path.read_text().splitlines()
    in_summary = False
    summary_data = []

    for line in lines:
        if line.strip() == '## Summarized Itinerary':
            in_summary = True
            summary_data.append(('📋 Summarized Itinerary', 'heading'))
            summary_data.append(('', 'blank'))
            continue

        if in_summary and line.startswith('## ') and 'Summarized' not in line:
            break

        if not in_summary:
            continue

        stripped = line.strip()

        if not stripped:
            summary_data.append(('', 'blank'))
        elif stripped.startswith('**Day') and stripped.endswith('**'):
            # Day header: **Day X | ...**
            text = stripped.strip('*')
            summary_data.append((text, 'day_header'))
        elif stripped.startswith('—') and 'Days' in stripped:
            # Section header: — VIENNA (Days 1-5) —
            summary_data.append((stripped, 'section_header'))
        elif re.match(r'^\d+\.', stripped):
            # Numbered item
            summary_data.append((stripped, 'item'))

    return summary_data


def insert_summary_into_docx(base_path: Path, summary_data: list, output_path: Path):
    """Insert summary paragraphs before 'Trip at a Glance' in the docx."""

    doc = Document(str(base_path))

    # Find insertion point
    target_element = None
    for para in doc.paragraphs:
        if '🗓️ Trip at a Glance' in para.text or 'Trip at a Glance' in para.text:
            target_element = para._element
            break

    if target_element is None:
        raise ValueError("Could not find 'Trip at a Glance' heading in the document")

    # Check if summary already exists (avoid duplicate insertion)
    for para in doc.paragraphs:
        if '📋 Summarized Itinerary' in para.text:
            print("Summary already exists in document. Skipping insertion.")
            doc.save(str(output_path))
            return

    # Insert paragraphs
    for text, ptype in summary_data:
        new_p = OxmlElement('w:p')

        if ptype == 'heading':
            pPr = OxmlElement('w:pPr')
            pStyle = OxmlElement('w:pStyle')
            pStyle.set(qn('w:val'), 'Heading3')
            pPr.append(pStyle)
            new_p.append(pPr)
            run = OxmlElement('w:r')
            t = OxmlElement('w:t')
            t.text = text
            t.set(qn('xml:space'), 'preserve')
            run.append(t)
            new_p.append(run)

        elif ptype == 'section_header':
            pPr = OxmlElement('w:pPr')
            jc = OxmlElement('w:jc')
            jc.set(qn('w:val'), 'center')
            pPr.append(jc)
            new_p.append(pPr)
            run = OxmlElement('w:r')
            rPr = OxmlElement('w:rPr')
            b = OxmlElement('w:b')
            rPr.append(b)
            i_el = OxmlElement('w:i')
            rPr.append(i_el)
            sz = OxmlElement('w:sz')
            sz.set(qn('w:val'), '22')
            rPr.append(sz)
            run.append(rPr)
            t = OxmlElement('w:t')
            t.text = text
            t.set(qn('xml:space'), 'preserve')
            run.append(t)
            new_p.append(run)

        elif ptype == 'day_header':
            run = OxmlElement('w:r')
            rPr = OxmlElement('w:rPr')
            b = OxmlElement('w:b')
            rPr.append(b)
            sz = OxmlElement('w:sz')
            sz.set(qn('w:val'), '22')
            rPr.append(sz)
            run.append(rPr)
            t = OxmlElement('w:t')
            t.text = text
            t.set(qn('xml:space'), 'preserve')
            run.append(t)
            new_p.append(run)

        elif ptype == 'item':
            pPr = OxmlElement('w:pPr')
            ind = OxmlElement('w:ind')
            ind.set(qn('w:left'), '360')
            pPr.append(ind)
            spacing = OxmlElement('w:spacing')
            spacing.set(qn('w:after'), '40')
            spacing.set(qn('w:line'), '260')
            pPr.append(spacing)
            new_p.append(pPr)
            run = OxmlElement('w:r')
            rPr = OxmlElement('w:rPr')
            sz = OxmlElement('w:sz')
            sz.set(qn('w:val'), '20')
            rPr.append(sz)
            run.append(rPr)
            t = OxmlElement('w:t')
            t.text = text
            t.set(qn('xml:space'), 'preserve')
            run.append(t)
            new_p.append(run)

        elif ptype == 'blank':
            pass

        target_element.addprevious(new_p)

    # Add separator before Trip at a Glance
    sep_p = OxmlElement('w:p')
    run = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = '—' * 50
    run.append(t)
    sep_p.append(run)
    target_element.addprevious(sep_p)

    doc.save(str(output_path))
    print(f"Summary inserted. Saved to {output_path}")


def main():
    parser = argparse.ArgumentParser(description="Insert summarized itinerary into docx")
    parser.add_argument("--base", required=True, help="Base .docx with rich formatting")
    parser.add_argument("--summary", required=True, help="Source .md file with summarized itinerary")
    parser.add_argument("--output", required=True, help="Output .docx path")
    args = parser.parse_args()

    summary_data = parse_summary_from_md(Path(args.summary))
    print(f"Parsed {len(summary_data)} summary lines from {args.summary}")

    insert_summary_into_docx(Path(args.base), summary_data, Path(args.output))


if __name__ == "__main__":
    main()
