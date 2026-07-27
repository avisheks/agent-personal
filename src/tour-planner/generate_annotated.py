#!/usr/bin/env python3
"""
Generate the annotated itinerary .docx from the richly-formatted annotated markdown.

Two modes of operation:

1. RESTYLE (--source-docx): Extract content from an existing annotated .docx and re-render
   it with updated styles from --reference-doc. Use when styles have changed but content is
   unchanged.

   python3 src/tour-planner/generate_annotated.py \
       --source-docx {TRIP-SLUG}/output/v2-trip-itinerary-annot-v3.docx \
       --reference-doc {TRIP-SLUG}/output/v2-trip-itinerary-latest.docx \
       --output {TRIP-SLUG}/output/v2-trip-itinerary-annot-latest.docx

2. MERGE (--base-md + --annotations-docx): Insert annotation blocks from a prior annotated
   version into an updated base markdown, then convert to .docx.

   python3 src/tour-planner/generate_annotated.py \
       --base-md {TRIP-SLUG}/output/v2-trip-itinerary-final.md \
       --annotations-docx {TRIP-SLUG}/output/v2-trip-itinerary-annot-v3.docx \
       --reference-doc {TRIP-SLUG}/output/v2-trip-itinerary-latest.docx \
       --output {TRIP-SLUG}/output/v2-trip-itinerary-annot-latest.docx

CRITICAL NOTES:
- The --base-md MUST be the richly-formatted v2-trip-itinerary-final.md (with emojis, route
  maps, full activity descriptions, proper paragraph-level formatting).
- NEVER use output/trip-itinerary-latest.md as --base-md. That file uses a simplified format
  (no emojis, markdown bullet-list activities, compact metadata) designed for a different
  pipeline (generate_itinerary.py → trip-itinerary-latest.docx).
- Using the simplified md produces: no emojis, no route maps, all activities as Compact style
  instead of Body Text, and missing content for later days.
"""

import argparse
import re
import subprocess
import tempfile
from pathlib import Path


# Annotation anchors: each annotation block is placed after the activity
# containing the keyword within the specified day range.
ANNOTATION_ANCHORS = [
    {'day': 2, 'keyword': 'Graben', 'keyword_alt': 'Kohlmarkt'},
    {'day': 2, 'keyword': 'Naturhistorisches', 'keyword_alt': 'NHM'},
    {'day': 3, 'keyword': 'Maze', 'keyword_alt': 'Schönbrunn Gardens'},
    {'day': 3, 'keyword': 'Wurstelprater', 'keyword_alt': 'Prater rides'},
    {'day': 4, 'keyword': 'Donaukanal', 'keyword_alt': 'canal stroll'},
    {'day': 5, 'keyword': 'Naschmarkt', 'keyword_alt': 'market'},
]


ACTIVITY_LINE_RE = re.compile(
    r'^(?:[🏛🎨🌳🚶🍽☕🎡🛍🏨🚇✈🚢🏰]️?\s*)?'
    r'(?:- \*\*)?'
    r'\d{1,2}:\d{2}\s*(?:AM|PM)'
)


def is_activity_line(line: str) -> bool:
    stripped = line.strip()
    return bool(stripped and ACTIVITY_LINE_RE.match(stripped))


def restyle(source_docx: Path, reference_doc: Path, output_docx: Path, output_md: Path,
            summary_md: Path | None = None):
    """Mode 1: Extract content from source docx and re-render with new styles.

    If --summary-md is provided, injects the Summarized Itinerary section from that file
    before the 'Trip at a Glance' heading (if not already present).

    IMPORTANT: The summarized itinerary source (--summary-md) should be the heat-optimized
    version (output/trip-itinerary-latest.md) which has already passed Layer 2 validation.
    On hot days (>30°C), outdoor activities must come first (before 11 AM) — never start a
    hot day with an indoor activity in the summary.
    """
    print(f"RESTYLE mode")
    print(f"  Source content: {source_docx}")
    print(f"  Style reference: {reference_doc}")

    # Extract content as GFM (preserves pipe tables, blockquotes, emojis)
    subprocess.run(
        ["pandoc", str(source_docx), "-o", str(output_md),
         "--from", "docx", "--to", "gfm", "--wrap=none"],
        capture_output=True, text=True, check=True,
    )
    print(f"  Extracted GFM: {output_md} ({sum(1 for _ in open(output_md))} lines)")

    # Inject summarized itinerary if provided and not already present
    if summary_md:
        _inject_summarized_itinerary(output_md, summary_md)

    # Re-render with reference-doc styles
    result = subprocess.run(
        ["pandoc", str(output_md), "-o", str(output_docx),
         "--from", "gfm", "--to", "docx", "--reference-doc", str(reference_doc)],
        capture_output=True, text=True,
    )
    if result.returncode != 0:
        print(f"ERROR: pandoc failed: {result.stderr}")
        raise RuntimeError("pandoc conversion failed")

    _verify_output(output_docx)


def _inject_summarized_itinerary(target_md: Path, summary_source_md: Path):
    """Insert the Summarized Itinerary section from summary_source_md into target_md.

    Inserts before the '### 🗓️ Trip at a Glance' heading. Skips if already present.
    The summary_source_md should contain a '## Summarized Itinerary' section with
    numbered lists per day.
    """
    target_lines = target_md.read_text().splitlines()

    # Check if already present
    if any('Summarized Itinerary' in line for line in target_lines):
        print(f"  Summarized Itinerary already present — skipping injection")
        return

    # Extract summarized itinerary from source
    source_lines = summary_source_md.read_text().splitlines()
    summary_start = None
    summary_end = None
    for i, line in enumerate(source_lines):
        if 'Summarized Itinerary' in line:
            summary_start = i
        elif summary_start and line.startswith('## ') and 'Summarized' not in line:
            summary_end = i
            break
        elif summary_start and line.strip() == '---':
            summary_end = i
            break

    if summary_start is None:
        print(f"  WARNING: No 'Summarized Itinerary' section found in {summary_source_md}")
        return

    summary_block = source_lines[summary_start:(summary_end or len(source_lines))]
    # Convert ## heading to ### for consistency with the annotated doc
    if summary_block and summary_block[0].startswith('## '):
        summary_block[0] = '#' + summary_block[0]  # ## -> ###

    # Find insertion point (before Trip at a Glance)
    insert_idx = None
    for i, line in enumerate(target_lines):
        if 'Trip at a Glance' in line:
            insert_idx = i
            break

    if insert_idx is None:
        print(f"  WARNING: Could not find 'Trip at a Glance' heading for insertion")
        return

    # Insert summary block before Trip at a Glance
    insertion = [''] + summary_block + ['']
    target_lines[insert_idx:insert_idx] = insertion
    target_md.write_text('\n'.join(target_lines))
    print(f"  Injected Summarized Itinerary ({len(summary_block)} lines) before Trip at a Glance")


def merge(base_md: Path, annotations_source: Path, reference_doc: Path,
          output_docx: Path, output_md: Path):
    """Mode 2: Merge annotations into base md, then convert to docx."""
    print(f"MERGE mode")
    print(f"  Base content: {base_md}")
    print(f"  Annotations from: {annotations_source}")
    print(f"  Style reference: {reference_doc}")

    # Extract annotations from source
    if annotations_source.suffix == '.docx':
        with tempfile.NamedTemporaryFile(suffix='.md', delete=False) as f:
            tmp = Path(f.name)
        subprocess.run(
            ["pandoc", str(annotations_source), "-o", str(tmp),
             "--from", "docx", "--to", "gfm", "--wrap=none"],
            capture_output=True, text=True, check=True,
        )
        annotations, cross_ref = _extract_annotations(tmp)
    else:
        annotations, cross_ref = _extract_annotations(annotations_source)

    print(f"  Found {len(annotations)} annotation blocks")
    if cross_ref:
        print(f"  Found cross-reference summary ({len(cross_ref)} chars)")

    # Merge into base
    merged = _merge_into_base(base_md, annotations, cross_ref)

    # Prevent overwriting source
    if base_md.resolve() == output_md.resolve():
        output_md = output_docx.with_name(output_docx.stem + '-merged.md')

    output_md.write_text(merged)
    print(f"  Merged .md: {output_md}")

    # Convert
    result = subprocess.run(
        ["pandoc", str(output_md), "-o", str(output_docx),
         "--from", "gfm", "--to", "docx", "--reference-doc", str(reference_doc)],
        capture_output=True, text=True,
    )
    if result.returncode != 0:
        print(f"ERROR: pandoc failed: {result.stderr}")
        raise RuntimeError("pandoc conversion failed")

    _verify_output(output_docx)


def _extract_annotations(md_path: Path) -> tuple[list[dict], str]:
    """Extract annotation blocks and cross-reference summary from markdown."""
    lines = md_path.read_text().splitlines()
    annotations = []
    cross_ref = ""

    # Cross-reference summary
    cr_start = None
    cr_end = None
    for i, line in enumerate(lines):
        if 'Cross-Reference Summary' in line:
            cr_start = i
        if cr_start and i > cr_start:
            if re.match(r'^#{1,3}\s+[✅📝]', line) or 'Packing Checklist' in line:
                cr_end = i
                break
    if cr_start:
        cross_ref = '\n'.join(lines[cr_start:(cr_end or len(lines))]).rstrip()

    # Annotation blocks
    current_day = 0
    anchor_counter = 0
    i = 0
    while i < len(lines):
        day_match = re.match(r'^#{1,3}\s+Day\s+(\d+)', lines[i])
        if day_match:
            current_day = int(day_match.group(1))

        if lines[i].strip().startswith('> 💡') or (lines[i].startswith('>') and '💡' in lines[i]):
            block_lines = []
            while i < len(lines):
                if lines[i].startswith('>'):
                    block_lines.append(lines[i])
                    i += 1
                elif lines[i].strip() == '' and i + 1 < len(lines) and lines[i + 1].startswith('>'):
                    block_lines.append(lines[i])
                    i += 1
                elif lines[i].strip() == '' and block_lines:
                    i += 1
                    break
                else:
                    break
            while block_lines and not block_lines[-1].strip():
                block_lines.pop()
            annotations.append({
                'block': '\n'.join(block_lines),
                'day': current_day,
                'anchor_idx': anchor_counter,
            })
            anchor_counter += 1
            continue
        i += 1

    return annotations, cross_ref


def _merge_into_base(base_md: Path, annotations: list[dict], cross_ref: str) -> str:
    """Merge annotation blocks into the base markdown at correct positions."""
    base_lines = base_md.read_text().splitlines()
    day_ranges = _get_day_ranges(base_lines)

    placed = []
    for ann in annotations:
        pos = _find_anchor(base_lines, ann, day_ranges)
        if pos is not None:
            placed.append((pos, ann['block']))
        else:
            print(f"  WARNING: Could not place (Day {ann['day']}): "
                  f"{ann['block'].split(chr(10))[0][:70]}...")

    # Insert in reverse order
    placed.sort(key=lambda x: x[0], reverse=True)
    for pos, block in placed:
        base_lines[pos + 1:pos + 1] = ['', block, '']

    result = '\n'.join(base_lines)
    if cross_ref:
        for marker in ['## Restaurant Summary', '### Restaurant Summary',
                       '## Packing', '### Packing', '## Rainy Day', '### Rainy Day']:
            if marker in result:
                result = result.replace(marker, cross_ref + '\n\n---\n\n' + marker)
                break
        else:
            result += '\n\n---\n\n' + cross_ref

    return result


def _get_day_ranges(lines: list[str]) -> dict[int, tuple[int, int]]:
    day_starts = []
    for i, line in enumerate(lines):
        m = re.match(r'^#{2,3}\s+Day\s+(\d+)', line)
        if m:
            day_starts.append((int(m.group(1)), i))
    ranges = {}
    for idx, (day_num, start) in enumerate(day_starts):
        end = day_starts[idx + 1][1] if idx + 1 < len(day_starts) else len(lines)
        ranges[day_num] = (start, end)
    return ranges


def _find_anchor(base_lines: list[str], annotation: dict, day_ranges: dict) -> int | None:
    day = annotation['day']
    if day not in day_ranges:
        return None
    start, end = day_ranges[day]
    idx = annotation['anchor_idx']
    if idx >= len(ANNOTATION_ANCHORS):
        return None
    anchor = ANNOTATION_ANCHORS[idx]

    keyword = anchor['keyword']
    keyword_alt = anchor.get('keyword_alt', '')

    best = None
    for i in range(start, min(end, len(base_lines))):
        line = base_lines[i]
        if keyword.lower() in line.lower() or (keyword_alt and keyword_alt.lower() in line.lower()):
            if is_activity_line(line):
                best = i
                break
            elif best is None:
                best = i

    if best is None:
        return None

    # Find end of activity block
    insert_after = best
    for j in range(best + 1, min(end, len(base_lines))):
        line = base_lines[j]
        stripped = line.strip()
        if not stripped or is_activity_line(line) or stripped.startswith('#') or stripped == '---':
            break
        insert_after = j
    return insert_after


def _verify_output(output_docx: Path):
    from docx import Document
    doc = Document(str(output_docx))
    blocks = [p for p in doc.paragraphs if p.style.name == 'Block Text']
    # Check for emoji presence in Body Text
    body_with_emoji = sum(1 for p in doc.paragraphs
                         if p.style.name == 'Body Text' and any(c in p.text for c in '🏛🎨🌳🚶🍽☕🎡🛍🏨🚇✈🚢🏰📍'))
    print(f"\n  Output: {output_docx}")
    print(f"  {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables")
    print(f"  {len(blocks)} suggestion blocks (Block Text)")
    print(f"  {body_with_emoji} activity lines with emojis (Body Text)")
    if body_with_emoji == 0:
        print(f"  ⚠️  WARNING: No emoji activity lines found! Did you use the wrong --base-md?")


def main():
    parser = argparse.ArgumentParser(
        description="Generate annotated itinerary .docx",
        epilog="Use --source-docx for restyle mode, or --base-md + --annotations-docx for merge mode.",
    )
    parser.add_argument("--source-docx", help="(Restyle) Existing annotated .docx to re-render with new styles")
    parser.add_argument("--base-md", help="(Merge) Richly-formatted base .md (v2-trip-itinerary-final.md)")
    parser.add_argument("--annotations-docx", help="(Merge) Prior annotated .docx to extract suggestion blocks from")
    parser.add_argument("--annotations-md", help="(Merge) Prior annotated .md (alternative to --annotations-docx)")
    parser.add_argument("--reference-doc", required=True, help="Styled .docx for pandoc --reference-doc (style source)")
    parser.add_argument("--output", required=True, help="Output .docx path")
    parser.add_argument("--output-md", help="Path for intermediate .md (default: <output>.md)")
    parser.add_argument("--summary-md", help="Source .md containing Summarized Itinerary to inject (e.g., output/trip-itinerary-latest.md)")
    args = parser.parse_args()

    output_docx = Path(args.output)
    output_md = Path(args.output_md) if args.output_md else output_docx.with_suffix('.md')
    ref_doc = Path(args.reference_doc)
    summary_md = Path(args.summary_md) if args.summary_md else None

    if not ref_doc.exists():
        raise FileNotFoundError(f"Reference .docx not found: {ref_doc}")
    if summary_md and not summary_md.exists():
        raise FileNotFoundError(f"Summary .md not found: {summary_md}")

    if args.source_docx:
        # RESTYLE mode
        source = Path(args.source_docx)
        if not source.exists():
            raise FileNotFoundError(f"Source .docx not found: {source}")
        restyle(source, ref_doc, output_docx, output_md, summary_md)

    elif args.base_md:
        # MERGE mode
        base = Path(args.base_md)
        if not base.exists():
            raise FileNotFoundError(f"Base .md not found: {base}")
        annot_src = Path(args.annotations_docx or args.annotations_md or '')
        if not annot_src.exists():
            raise FileNotFoundError(f"Annotations source not found: {annot_src}")
        merge(base, annot_src, ref_doc, output_docx, output_md)

    else:
        parser.error("Must provide either --source-docx (restyle) or --base-md (merge)")


if __name__ == "__main__":
    main()
