#!/usr/bin/env python3
"""
Validate computed PnL against a reference spreadsheet.
Compares trade-level and monthly-level totals, reports discrepancies.

Usage:
    python3 src/validate_pnl.py \
        --computed .local/options-pnl/out/computed-pnl.json \
        --reference ".local/options-pnl/ref/[Options] Ultimate Options Tracking Spreadsheet - Version 2 - Q12026.xlsx"
"""

import argparse
import json
from pathlib import Path
from datetime import datetime
from collections import defaultdict

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl required. Install with: pip3 install openpyxl")
    exit(1)


MONTH_MAP = {
    "Jan": "2026-01", "Feb": "2026-02", "Mar": "2026-03",
    "Apr": "2026-04", "May": "2026-05", "Jun": "2026-06",
    "Jul": "2026-07", "Aug": "2026-08", "Sep": "2026-09",
    "Oct": "2026-10", "Nov": "2026-11", "Dec": "2026-12",
}

BROKER_MAP = {
    "Fidelity": "Fidelity",
    "ToS": "Thinkorswim",
    "Tasty": "Tastytrade",
}


def parse_reference(filepath: Path) -> dict:
    """Parse reference xlsx into structured data."""
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    trades = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        month_raw = row[0]
        event = row[7]
        pnl = row[14]

        if month_raw is None or event is None or pnl is None:
            continue

        event_clean = str(event).strip()
        if event_clean not in ("CC", "CSP"):
            continue

        month = MONTH_MAP.get(str(month_raw), str(month_raw))
        broker_raw = str(row[1]).strip()
        broker = BROKER_MAP.get(broker_raw, broker_raw)

        opn = row[3]
        cls = row[5]
        ticker = str(row[6]).strip() if row[6] else ""
        strike = row[8]
        qty = row[9]
        fill_price = row[10]
        close_price = row[11]

        trades.append({
            "month": month,
            "broker": broker,
            "strategy": event_clean,
            "underlying": ticker,
            "strike": float(strike) if strike else 0,
            "qty": int(qty) if qty else 0,
            "open_price": float(fill_price) if fill_price else 0,
            "close_price": float(close_price) if close_price else 0,
            "pnl": float(pnl),
            "date_opened": opn.strftime("%Y-%m-%d") if opn else None,
            "date_closed": cls.strftime("%Y-%m-%d") if cls else None,
        })

    wb.close()
    return trades


def load_computed(filepath: Path, ref_months: set) -> list:
    """Load computed trades, filtered to reference months only."""
    with open(filepath) as f:
        data = json.load(f)

    trades = []
    for t in data["closed_trades"]:
        month = t["date_closed"][:7] if t["date_closed"] else None
        if month in ref_months:
            trades.append(t)
    return trades


def match_key(t: dict) -> tuple:
    """Generate a matching key for a trade."""
    return (
        t["broker"],
        t["underlying"],
        t["strategy"],
        float(t["strike"]),
        int(t["qty"]),
        t.get("date_closed"),
    )


def fuzzy_match_key(t: dict) -> tuple:
    """Looser matching key (no close date, for finding near-matches)."""
    return (
        t["broker"],
        t["underlying"],
        t["strategy"],
        float(t["strike"]),
        int(t["qty"]),
    )


def validate(ref_trades: list, computed_trades: list, tolerance: float = 5.0) -> dict:
    """
    Compare reference and computed trades.
    Returns a validation report dict.
    """
    report = {
        "status": "PASS",
        "summary": {},
        "monthly_comparison": [],
        "matched_trades": [],
        "mismatched_pnl": [],
        "missing_from_computed": [],
        "extra_in_computed": [],
    }

    # Monthly totals comparison
    ref_by_month = defaultdict(float)
    comp_by_month = defaultdict(float)
    ref_count_month = defaultdict(int)
    comp_count_month = defaultdict(int)

    for t in ref_trades:
        ref_by_month[t["month"]] += t["pnl"]
        ref_count_month[t["month"]] += 1
    for t in computed_trades:
        m = t["date_closed"][:7]
        comp_by_month[m] += t["net_pnl"]
        comp_count_month[m] += 1

    all_months = sorted(set(ref_by_month.keys()) | set(comp_by_month.keys()))
    for month in all_months:
        ref_pnl = round(ref_by_month.get(month, 0), 2)
        comp_pnl = round(comp_by_month.get(month, 0), 2)
        diff = round(comp_pnl - ref_pnl, 2)
        report["monthly_comparison"].append({
            "month": month,
            "ref_trades": ref_count_month.get(month, 0),
            "computed_trades": comp_count_month.get(month, 0),
            "ref_pnl": ref_pnl,
            "computed_pnl": comp_pnl,
            "difference": diff,
            "status": "OK" if abs(diff) <= tolerance else "MISMATCH",
        })

    # Broker totals
    ref_by_broker = defaultdict(float)
    comp_by_broker = defaultdict(float)
    for t in ref_trades:
        ref_by_broker[t["broker"]] += t["pnl"]
    for t in computed_trades:
        comp_by_broker[t["broker"]] += t["net_pnl"]

    # Trade-level matching
    ref_remaining = list(ref_trades)
    comp_remaining = list(computed_trades)

    matched_ref = set()
    matched_comp = set()

    # Pass 1: exact match on key + close date
    for i, rt in enumerate(ref_remaining):
        rkey = match_key(rt)
        for j, ct in enumerate(comp_remaining):
            if j in matched_comp:
                continue
            ckey = match_key(ct)
            if rkey == ckey:
                pnl_diff = abs(rt["pnl"] - ct["net_pnl"])
                if pnl_diff <= tolerance:
                    report["matched_trades"].append({
                        "ref": rt, "computed": ct, "pnl_diff": round(pnl_diff, 2)
                    })
                else:
                    report["mismatched_pnl"].append({
                        "ref": rt, "computed": ct,
                        "ref_pnl": round(rt["pnl"], 2),
                        "computed_pnl": round(ct["net_pnl"], 2),
                        "difference": round(ct["net_pnl"] - rt["pnl"], 2),
                    })
                matched_ref.add(i)
                matched_comp.add(j)
                break

    # Pass 2: fuzzy match (no close date) for remaining
    for i, rt in enumerate(ref_remaining):
        if i in matched_ref:
            continue
        rkey = fuzzy_match_key(rt)
        for j, ct in enumerate(comp_remaining):
            if j in matched_comp:
                continue
            ckey = fuzzy_match_key(ct)
            if rkey == ckey:
                pnl_diff = abs(rt["pnl"] - ct["net_pnl"])
                if pnl_diff <= tolerance:
                    report["matched_trades"].append({
                        "ref": rt, "computed": ct, "pnl_diff": round(pnl_diff, 2),
                        "note": "fuzzy match (date mismatch)"
                    })
                else:
                    report["mismatched_pnl"].append({
                        "ref": rt, "computed": ct,
                        "ref_pnl": round(rt["pnl"], 2),
                        "computed_pnl": round(ct["net_pnl"], 2),
                        "difference": round(ct["net_pnl"] - rt["pnl"], 2),
                        "note": "fuzzy match (date mismatch)"
                    })
                matched_ref.add(i)
                matched_comp.add(j)
                break

    # Unmatched
    for i, rt in enumerate(ref_remaining):
        if i not in matched_ref:
            report["missing_from_computed"].append(rt)

    for j, ct in enumerate(comp_remaining):
        if j not in matched_comp:
            report["extra_in_computed"].append(ct)

    # Summary
    total_ref = round(sum(t["pnl"] for t in ref_trades), 2)
    total_comp = round(sum(t["net_pnl"] for t in computed_trades), 2)

    report["summary"] = {
        "ref_total_pnl": total_ref,
        "computed_total_pnl": total_comp,
        "total_difference": round(total_comp - total_ref, 2),
        "ref_trade_count": len(ref_trades),
        "computed_trade_count": len(computed_trades),
        "matched": len(report["matched_trades"]),
        "pnl_mismatches": len(report["mismatched_pnl"]),
        "missing_from_computed": len(report["missing_from_computed"]),
        "extra_in_computed": len(report["extra_in_computed"]),
        "ref_by_broker": {k: round(v, 2) for k, v in ref_by_broker.items()},
        "computed_by_broker": {k: round(v, 2) for k, v in comp_by_broker.items()},
    }

    # Set overall status
    if report["mismatched_pnl"] or report["missing_from_computed"] or report["extra_in_computed"]:
        report["status"] = "FAIL"
    if abs(total_comp - total_ref) > tolerance * len(ref_trades):
        report["status"] = "FAIL"

    return report


def print_report(report: dict):
    """Print human-readable validation report."""
    s = report["summary"]
    status_icon = "PASS" if report["status"] == "PASS" else "FAIL"

    print(f"\n{'='*60}")
    print(f"  VALIDATION RESULT: {status_icon}")
    print(f"{'='*60}\n")

    print(f"  Reference PnL:  ${s['ref_total_pnl']:>10,.2f}  ({s['ref_trade_count']} trades)")
    print(f"  Computed PnL:   ${s['computed_total_pnl']:>10,.2f}  ({s['computed_trade_count']} trades)")
    print(f"  Difference:     ${s['total_difference']:>10,.2f}")
    print()

    print("  Monthly Comparison:")
    print(f"  {'Month':<10} {'Ref Trades':>10} {'Comp Trades':>11} {'Ref PnL':>10} {'Comp PnL':>10} {'Diff':>10} {'Status'}")
    print(f"  {'-'*10} {'-'*10} {'-'*11} {'-'*10} {'-'*10} {'-'*10} {'-'*8}")
    for m in report["monthly_comparison"]:
        print(f"  {m['month']:<10} {m['ref_trades']:>10} {m['computed_trades']:>11} ${m['ref_pnl']:>9,.2f} ${m['computed_pnl']:>9,.2f} ${m['difference']:>9,.2f} {m['status']}")
    print()

    print("  Broker Comparison:")
    all_brokers = sorted(set(list(s["ref_by_broker"].keys()) + list(s["computed_by_broker"].keys())))
    print(f"  {'Broker':<15} {'Ref PnL':>10} {'Comp PnL':>10} {'Diff':>10}")
    print(f"  {'-'*15} {'-'*10} {'-'*10} {'-'*10}")
    for b in all_brokers:
        rp = s["ref_by_broker"].get(b, 0)
        cp = s["computed_by_broker"].get(b, 0)
        print(f"  {b:<15} ${rp:>9,.2f} ${cp:>9,.2f} ${cp-rp:>9,.2f}")
    print()

    print(f"  Trade Matching:")
    print(f"    Matched (PnL within tolerance): {s['matched']}")
    print(f"    PnL mismatches:                 {s['pnl_mismatches']}")
    print(f"    Missing from computed:          {s['missing_from_computed']}")
    print(f"    Extra in computed:              {s['extra_in_computed']}")
    print()

    if report["mismatched_pnl"]:
        print("  PnL Mismatches (trades found but PnL differs):")
        print(f"  {'Broker':<12} {'Ticker':<8} {'Strategy':<5} {'Strike':>7} {'Ref PnL':>9} {'Comp PnL':>9} {'Diff':>8}")
        print(f"  {'-'*12} {'-'*8} {'-'*5} {'-'*7} {'-'*9} {'-'*9} {'-'*8}")
        for m in report["mismatched_pnl"]:
            r = m["ref"]
            print(f"  {r['broker']:<12} {r['underlying']:<8} {r['strategy']:<5} ${r['strike']:>6.0f} ${m['ref_pnl']:>8,.2f} ${m['computed_pnl']:>8,.2f} ${m['difference']:>7,.2f}")
        print()

    if report["missing_from_computed"]:
        print("  Missing from Computed (in reference but not found):")
        print(f"  {'Broker':<12} {'Ticker':<8} {'Strategy':<5} {'Strike':>7} {'Qty':>4} {'Close Date':<12} {'Ref PnL':>9}")
        print(f"  {'-'*12} {'-'*8} {'-'*5} {'-'*7} {'-'*4} {'-'*12} {'-'*9}")
        for t in report["missing_from_computed"]:
            print(f"  {t['broker']:<12} {t['underlying']:<8} {t['strategy']:<5} ${t['strike']:>6.0f} {t['qty']:>4} {t.get('date_closed','?'):<12} ${t['pnl']:>8,.2f}")
        print()

    if report["extra_in_computed"]:
        print("  Extra in Computed (not in reference):")
        print(f"  {'Broker':<12} {'Ticker':<8} {'Strategy':<5} {'Strike':>7} {'Qty':>4} {'Close Date':<12} {'Net PnL':>9}")
        print(f"  {'-'*12} {'-'*8} {'-'*5} {'-'*7} {'-'*4} {'-'*12} {'-'*9}")
        for t in report["extra_in_computed"][:20]:
            print(f"  {t['broker']:<12} {t['underlying']:<8} {t['strategy']:<5} ${t['strike']:>6.0f} {t['qty']:>4} {t.get('date_closed','?'):<12} ${t['net_pnl']:>8,.2f}")
        if len(report["extra_in_computed"]) > 20:
            print(f"  ... and {len(report['extra_in_computed']) - 20} more")
        print()


def main():
    parser = argparse.ArgumentParser(description="Validate computed PnL against reference")
    parser.add_argument("--computed", required=True, help="Path to computed-pnl.json")
    parser.add_argument("--reference", required=True, help="Path to reference xlsx")
    parser.add_argument("--tolerance", type=float, default=5.0, help="Per-trade PnL tolerance in dollars (default: $5)")
    parser.add_argument("--output", help="Optional: write validation report as JSON")
    args = parser.parse_args()

    ref_trades = parse_reference(Path(args.reference))
    ref_months = set(t["month"] for t in ref_trades)

    computed_trades = load_computed(Path(args.computed), ref_months)

    report = validate(ref_trades, computed_trades, tolerance=args.tolerance)
    print_report(report)

    if args.output:
        out_path = Path(args.output)
        with open(out_path, "w") as f:
            json.dump(report, f, indent=2, default=str)
        print(f"  Validation report written to: {out_path}")

    exit(0 if report["status"] == "PASS" else 1)


if __name__ == "__main__":
    main()
