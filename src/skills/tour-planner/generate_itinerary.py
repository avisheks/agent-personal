#!/usr/bin/env python3
"""
Orchestrator: generates the final trip-itinerary-latest.docx from a styled base + aligned .md.

Uses pandoc with --reference-doc to inherit styles from the base .docx while using
the aligned .md as the content source. This preserves heading styles, paragraph spacing,
font choices, and table formatting from the base document.

Usage:
    python3 src/tour-planner/generate_itinerary.py \
        --base output/trip-itinerary-v2.docx \
        --md output/trip-itinerary-latest.md \
        --output output/trip-itinerary-latest.docx
"""

import argparse
import subprocess
from pathlib import Path


def main():
    parser = argparse.ArgumentParser(description="Generate final itinerary .docx")
    parser.add_argument("--base", required=True, help="Styled base .docx (provides style/formatting reference)")
    parser.add_argument("--md", required=True, help="Aligned .md with correct content (source of truth)")
    parser.add_argument("--output", required=True, help="Output .docx path")
    args = parser.parse_args()

    base_path = Path(args.base)
    md_path = Path(args.md)
    output_path = Path(args.output)

    if not base_path.exists():
        raise FileNotFoundError(f"Base .docx not found: {base_path}")
    if not md_path.exists():
        raise FileNotFoundError(f"Source .md not found: {md_path}")

    # Generate .docx using pandoc with reference-doc for style inheritance
    print(f"Generating {output_path}")
    print(f"  Content from: {md_path}")
    print(f"  Styles from:  {base_path}")

    result = subprocess.run(
        [
            "pandoc", str(md_path),
            "-o", str(output_path),
            "--from", "markdown",
            "--to", "docx",
            "--reference-doc", str(base_path),
        ],
        capture_output=True,
        text=True,
    )

    if result.returncode != 0:
        print(f"ERROR: pandoc failed with code {result.returncode}")
        print(f"  stderr: {result.stderr}")
        raise RuntimeError("pandoc conversion failed")

    # Verify output
    from docx import Document
    doc = Document(str(output_path))
    hotel_returns = [p.text for p in doc.paragraphs if 'Return to hotel' in p.text]
    if hotel_returns:
        print(f"  WARNING: {len(hotel_returns)} 'Return to hotel' entries found in output!")

    print(f"\nDone. {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables.")
    print(f"Output: {output_path}")


if __name__ == "__main__":
    main()
